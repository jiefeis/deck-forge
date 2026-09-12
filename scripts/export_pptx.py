#!/usr/bin/env python3
"""Turn a generated deck-forge HTML deck into a native, editable PowerPoint file.

    python export_pptx.py build <deck>/index.html <out.pptx> [--font NAME]
                                [--image-only] [--browser-executable PATH]
    python export_pptx.py diff  <ref_png_dir> <pptx_render_dir> [--erode N] [--json]

`build` reads geometry back from the browser instead of re-implementing CSS:
the deck is opened in Playwright, entry animations are frozen, and every
visual element is emitted as a PPTX shape at its measured box — text boxes
with their inline bold/colour runs, autoshapes for boxes and CSS triangles,
real tables, and pictures for <svg>/<img>. `--image-only` instead places one
full-bleed screenshot per slide (pixel-identical, nothing editable).

`diff` compares the source deck's own page renders (`export_pdf.py
--keep-pngs`) with the PPTX renders (`render_pptx.ps1`) and reports, per page,
how much of the difference survives an erosion pass. Thin differences are
glyph rasterisation; surviving solid blocks are layout defects.

Both commands are standalone: paths are arguments, cwd does not matter.
"""
from __future__ import annotations

import argparse
import json
import os
import sys
import tempfile
from pathlib import Path

# ---------------------------------------------------------------------------
# Units. The deck stage is 1920x1080 px and PowerPoint's 16:9 canvas is
# 13.333x7.5 in, so 1 px = 6350 EMU and (at 144 px/in) 1 px = 0.5 pt.
# ---------------------------------------------------------------------------
STAGE_W, STAGE_H = 1920.0, 1080.0
EMU_PER_PX = 6350
NS_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"

FREEZE_CSS = """
*{transition:none !important;animation:none !important;}
.reveal{opacity:1 !important;transform:none !important;}
"""

# Walk one slide (index passed in) and describe every visual element.
WALK_JS = r"""
(si) => {
  const INLINE = new Set(['B','SPAN','I','A','EM','STRONG','BR','SUP','SUB','U','SMALL','CODE']);
  const rgb = s => {
    const m = (s || '').match(/rgba?\(([\d.]+),\s*([\d.]+),\s*([\d.]+)(?:,\s*([\d.]+))?\)/);
    if (!m) return null;
    if (m[4] !== undefined && parseFloat(m[4]) === 0) return null;
    return [Math.round(+m[1]), Math.round(+m[2]), Math.round(+m[3])];
  };
  const slide = document.querySelectorAll('section.slide, .slide')[si];
  const sb = slide.getBoundingClientRect();
  const R = el => { const r = el.getBoundingClientRect();
    return {x: r.left - sb.left, y: r.top - sb.top, w: r.width, h: r.height}; };
  const items = [];
  let rasterN = 0;

  // Pseudo-elements never appear in the DOM walk: the template's edge bar,
  // cover rules and list bullets live here.
  for (const which of ['::before', '::after']) {
    const ps = getComputedStyle(slide, which);
    if (!ps || ps.content === 'none' || ps.content === 'normal') continue;
    const w = parseFloat(ps.width) || 0, h = parseFloat(ps.height) || 0;
    if (w <= 0 || h <= 0) continue;
    const x = ps.right !== 'auto' && ps.left === 'auto' ? sb.width - w - (parseFloat(ps.right) || 0)
                                                        : (parseFloat(ps.left) || 0);
    const y = ps.bottom !== 'auto' && ps.top === 'auto' ? sb.height - h - (parseFloat(ps.bottom) || 0)
                                                        : (parseFloat(ps.top) || 0);
    const fill = rgb(ps.backgroundColor);
    if (fill) items.push({kind: 'rect', box: {x, y, w, h}, fill, border: null, radius: 0});
  }

  const runsOf = el => {
    const runs = [];
    const walk = (n, bold, color) => {
      n.childNodes.forEach(c => {
        if (c.nodeType === 3) { if (c.nodeValue) runs.push({t: c.nodeValue, b: bold, c: color}); }
        else if (c.nodeType === 1) {
          if (c.tagName === 'BR') { runs.push({t: '\n', b: bold, c: color}); return; }
          const cs = getComputedStyle(c);
          walk(c, bold || +cs.fontWeight >= 600, rgb(cs.color) || color);
        }
      });
    };
    const cs = getComputedStyle(el);
    walk(el, +cs.fontWeight >= 600, rgb(cs.color));
    return runs.filter(r => r.t.trim() !== '' || r.t === '\n');
  };

  const tableOf = (t, box) => {
    const rows = [];
    t.querySelectorAll('tr').forEach(tr => {
      const cells = [];
      tr.querySelectorAll('th,td').forEach(td => {
        const cs = getComputedStyle(td), r = td.getBoundingClientRect();
        cells.push({runs: runsOf(td), align: cs.textAlign, size: parseFloat(cs.fontSize),
                    fill: rgb(cs.backgroundColor), bold: +cs.fontWeight >= 600,
                    w: r.width, h: r.height, border: rgb(cs.borderTopColor),
                    bw: parseFloat(cs.borderTopWidth) || 0});
      });
      rows.push(cells);
    });
    return {kind: 'table', box, rows};
  };

  // A CSS-border triangle: zero content box, one opaque border side.
  // clientWidth/Height exclude borders, so this holds under box-sizing:border-box
  // too (where computed width reports the clamped border-box, not the 0 content).
  const triangleOf = (el, cs, box) => {
    if (el.clientWidth > 1 || el.clientHeight > 1 || el.children.length || (el.textContent || '').trim()) return null;
    const sides = {right: cs.borderLeftColor, left: cs.borderRightColor,
                   down: cs.borderTopColor, up: cs.borderBottomColor};
    const widths = {right: cs.borderLeftWidth, left: cs.borderRightWidth,
                    down: cs.borderTopWidth, up: cs.borderBottomWidth};
    const opaque = Object.keys(sides).filter(k => rgb(sides[k]) && parseFloat(widths[k]) > 0);
    if (opaque.length !== 1 || box.w < 2 || box.h < 2) return null;
    return {kind: 'tri', dir: opaque[0], box, fill: rgb(sides[opaque[0]])};
  };

  const visit = el => {
    const cs = getComputedStyle(el);
    if (cs.display === 'none' || cs.visibility === 'hidden' || parseFloat(cs.opacity) === 0) return;
    const box = R(el);
    const tag = el.tagName.toUpperCase();
    if (tag === 'SCRIPT' || tag === 'STYLE') return;
    if (tag === 'TABLE') { items.push(tableOf(el, box)); return; }
    if (tag === 'SVG' || tag === 'IMG' || tag === 'CANVAS' || tag === 'VIDEO') {
      el.setAttribute('data-x2p', String(rasterN));
      items.push({kind: 'raster', box, ref: rasterN++});
      return;
    }
    const tri = triangleOf(el, cs, box);
    if (tri) { items.push(tri); return; }

    const fill = rgb(cs.backgroundColor);
    const bw = parseFloat(cs.borderTopWidth) || 0, bc = rgb(cs.borderTopColor);
    const rad = parseFloat(cs.borderTopLeftRadius) || 0;
    if ((fill || (bw > 0 && bc)) && box.w > 0.5 && box.h > 0.5)
      items.push({kind: 'rect', box, fill, radius: rad,
                  border: bw > 0 && bc ? {w: bw, c: bc} : null});

    const kids = [...el.children];
    const blockKids = kids.filter(k => !INLINE.has(k.tagName.toUpperCase()));
    const isTextLeaf = el.innerText && el.innerText.trim() && blockKids.length === 0;
    if (isTextLeaf) {
      let pre = '';
      const bf = getComputedStyle(el, '::before');
      if (bf && bf.content && !['none', 'normal', '""', "''"].includes(bf.content))
        pre = bf.content.replace(/^["']|["']$/g, '');
      const pl = parseFloat(cs.paddingLeft) || 0, pr = parseFloat(cs.paddingRight) || 0;
      const pt = parseFloat(cs.paddingTop) || 0, pb = parseFloat(cs.paddingBottom) || 0;
      items.push({kind: 'text', box,
        inner: {x: box.x + pl, y: box.y + pt, w: box.w - pl - pr, h: box.h - pt - pb},
        size: parseFloat(cs.fontSize),
        lh: parseFloat(cs.lineHeight) || parseFloat(cs.fontSize) * 1.4,
        align: cs.textAlign, color: rgb(cs.color), ls: parseFloat(cs.letterSpacing) || 0,
        runs: runsOf(el), pre});
      return;
    }
    kids.forEach(visit);
  };
  [...slide.children].forEach(visit);
  return items;
}
"""


# ---------------------------------------------------------------------------
# build
# ---------------------------------------------------------------------------
def _launch(pw, exe):
    kw = {}
    if exe:
        kw["executable_path"] = exe
    elif os.environ.get("DECK_FORGE_BROWSER_EXECUTABLE"):
        kw["executable_path"] = os.environ["DECK_FORGE_BROWSER_EXECUTABLE"]
    return pw.chromium.launch(**kw)


def _show_only(page, i):
    page.evaluate(
        "(i)=>document.querySelectorAll('section.slide, .slide').forEach((s,k)=>{"
        "s.classList.toggle('active',k===i);s.classList.toggle('visible',k===i);})", i)
    page.wait_for_timeout(120)


def extract(html: Path, exe: str | None, image_only: bool, scratch: Path):
    from playwright.sync_api import sync_playwright
    slides = []
    with sync_playwright() as pw:
        browser = _launch(pw, exe)
        page = browser.new_page(viewport={"width": int(STAGE_W), "height": int(STAGE_H)},
                                device_scale_factor=2)
        page.goto(html.resolve().as_uri())
        page.wait_for_timeout(800)
        page.add_style_tag(content=FREEZE_CSS)
        n = page.evaluate("document.querySelectorAll('section.slide, .slide').length")
        if n == 0:
            raise SystemExit("no .slide sections found - is this a deck-forge HTML deck?")
        for i in range(n):
            _show_only(page, i)
            if image_only:
                png = scratch / f"slide-{i + 1:03d}.png"
                page.screenshot(path=str(png), clip={"x": 0, "y": 0, "width": STAGE_W, "height": STAGE_H})
                slides.append({"items": [{"kind": "raster", "png": str(png),
                                          "box": {"x": 0, "y": 0, "w": STAGE_W, "h": STAGE_H}}]})
                continue
            items = page.evaluate(WALK_JS, i)
            for it in items:
                if it["kind"] == "raster":
                    png = scratch / f"s{i + 1:03d}-r{it['ref']}.png"
                    page.locator(f'[data-x2p="{it["ref"]}"]').screenshot(path=str(png), omit_background=True)
                    it["png"] = str(png)
            page.evaluate("document.querySelectorAll('[data-x2p]').forEach(e=>e.removeAttribute('data-x2p'))")
            slides.append({"items": items})
        browser.close()
    return slides


def build_pptx(slides, out: Path, font: str):
    from lxml import etree
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Emu, Pt

    def E(px): return Emu(int(round(px * EMU_PER_PX)))
    def P(px): return Pt(px * 0.5)
    ALIGN = {"left": PP_ALIGN.LEFT, "right": PP_ALIGN.RIGHT, "center": PP_ALIGN.CENTER,
             "justify": PP_ALIGN.JUSTIFY, "start": PP_ALIGN.LEFT, "end": PP_ALIGN.RIGHT}

    def clamp(x, y, w, h):
        # Slack added so PowerPoint does not re-wrap can push full-width boxes off
        # the canvas. Invisible when presenting; in edit view the frames show as
        # stray outlines outside the slide, which users report as overflow.
        x2, y2 = min(STAGE_W, x + w), min(STAGE_H, y + h)
        x, y = max(0.0, x), max(0.0, y)
        return x, y, max(1.0, x2 - x), max(1.0, y2 - y)

    def set_font(run):
        run.font.name = font
        rPr = run._r.get_or_add_rPr()
        # Without a CJK lang PowerPoint applies Latin line-breaking to Chinese:
        # 禁则 fails and 、，。」 land at line starts. eaLnBrk alone does nothing.
        rPr.set("lang", "zh-CN"); rPr.set("altLang", "en-US")
        for tag in ("ea", "cs"):          # latin alone leaves CJK to the theme font
            el = rPr.find(NS_A + tag)
            if el is None:
                el = etree.SubElement(rPr, NS_A + tag)
            el.set("typeface", font)

    def no_autofit(tf):                   # else PowerPoint silently shrinks text
        bodyPr = tf._txBody.find(NS_A + "bodyPr")
        for t in ("normAutofit", "spAutoFit"):
            e = bodyPr.find(NS_A + t)
            if e is not None:
                bodyPr.remove(e)
        etree.SubElement(bodyPr, NS_A + "noAutofit")

    def para_props(p, align, lh_px):
        p.alignment = ALIGN.get(align, PP_ALIGN.LEFT)
        pf = p._pPr if p._pPr is not None else p._p.get_or_add_pPr()
        pf.set("marL", "0"); pf.set("indent", "0")
        pf.set("eaLnBrk", "1"); pf.set("hangingPunct", "1"); pf.set("latinLnBrk", "0")
        ln = etree.SubElement(pf, NS_A + "lnSpc")
        etree.SubElement(ln, NS_A + "spcPts").set("val", str(int(round(lh_px * 0.5 * 100))))
        for tag in ("spcBef", "spcAft"):
            etree.SubElement(etree.SubElement(pf, NS_A + tag), NS_A + "spcPts").set("val", "0")

    def add_runs(p, runs, size_px, default_color, ls_px=0.0, force_bold=False):
        for r in runs:
            run = p.add_run()
            run.text = r["t"]
            run.font.size = P(size_px)
            run.font.bold = bool(r["b"]) or force_bold
            run.font.color.rgb = RGBColor(*(r["c"] or default_color or [0, 0, 0]))
            set_font(run)
            if ls_px:
                run._r.get_or_add_rPr().set("spc", str(int(round(ls_px * 50))))

    def add_text(slide, it):
        b = it.get("inner") or it["box"]
        x, y, w, h = clamp(b["x"] - 2, b["y"] - 2, b["w"] + 8, b["h"] + 8)
        tb = slide.shapes.add_textbox(E(x), E(y), E(w), E(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.TOP
        no_autofit(tf)
        paras, cur = [], []
        for r in it["runs"]:
            if r["t"] == "\n":
                paras.append(cur); cur = []
            else:
                cur.append(r)
        paras.append(cur)
        if it.get("pre") and paras and paras[0]:
            paras[0] = [{"t": it["pre"], "b": False, "c": it.get("color")}] + paras[0]
        for k, pr in enumerate(paras):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            para_props(p, it.get("align", "left"), it["lh"])
            add_runs(p, pr, it["size"], it.get("color"), it.get("ls", 0))

    def add_rect(slide, it):
        b = it["box"]
        rad = it.get("radius") or 0
        x, y, w, h = clamp(b["x"], b["y"], b["w"], b["h"])
        sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rad > 1 else MSO_SHAPE.RECTANGLE,
                                    E(x), E(y), E(w), E(h))
        if rad > 1:
            sh.adjustments[0] = min(0.5, rad / max(1.0, min(w, h)))
        if it.get("fill"):
            sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor(*it["fill"])
        else:
            sh.fill.background()
        if it.get("border"):
            sh.line.color.rgb = RGBColor(*it["border"]["c"])
            sh.line.width = Pt(max(0.5, it["border"]["w"] * 0.5))
        else:
            sh.line.fill.background()
        sh.shadow.inherit = False

    def add_tri(slide, it):
        b = it["box"]
        cx, cy = b["x"] + b["w"] / 2, b["y"] + b["h"] / 2
        rot = {"up": 0, "right": 90, "down": 180, "left": 270}[it["dir"]]
        w, h = (b["h"], b["w"]) if rot in (90, 270) else (b["w"], b["h"])
        sh = slide.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE,
                                    E(cx - w / 2), E(cy - h / 2), E(w), E(h))
        sh.rotation = rot
        sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor(*(it["fill"] or [128, 128, 128]))
        sh.line.fill.background()
        sh.shadow.inherit = False

    def add_table(slide, it):
        b, rows = it["box"], it["rows"]
        ncol = max(len(r) for r in rows)
        x, y, w, h = clamp(b["x"], b["y"], b["w"], b["h"])
        tbl = slide.shapes.add_table(len(rows), ncol, E(x), E(y), E(w), E(h)).table
        tbl.first_row = False
        for ci in range(ncol):
            tbl.columns[ci].width = E(rows[0][ci]["w"] if ci < len(rows[0]) else w / ncol)
        for ri, row in enumerate(rows):
            tbl.rows[ri].height = E(row[0]["h"])
            for ci, cell in enumerate(row):
                c = tbl.cell(ri, ci)
                c.margin_left = c.margin_right = E(18)
                c.margin_top = c.margin_bottom = E(10)
                c.vertical_anchor = MSO_ANCHOR.MIDDLE
                c.fill.solid(); c.fill.fore_color.rgb = RGBColor(*(cell.get("fill") or [255, 255, 255]))
                tf = c.text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                first = True
                for r in cell["runs"]:
                    if r["t"] == "\n":
                        p = tf.add_paragraph(); continue
                    if first:
                        p.alignment = ALIGN.get(cell.get("align", "left"), PP_ALIGN.LEFT); first = False
                    add_runs(p, [r], cell["size"], None, force_bold=bool(cell.get("bold")))

    def add_raster(slide, it):
        b = it["box"]
        x, y, w, h = clamp(b["x"], b["y"], b["w"], b["h"])
        slide.shapes.add_picture(it["png"], E(x), E(y), E(w), E(h))

    prs = Presentation()
    prs.slide_width, prs.slide_height = E(STAGE_W), E(STAGE_H)
    blank = prs.slide_layouts[6]
    handlers = {"text": add_text, "rect": add_rect, "tri": add_tri,
                "table": add_table, "raster": add_raster}
    for s in slides:
        slide = prs.slides.add_slide(blank)
        for it in s["items"]:
            handlers[it["kind"]](slide, it)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    return len(slides)


def cmd_build(a):
    html = Path(a.html)
    if not html.is_file():
        raise SystemExit(f"not found: {html}")
    with tempfile.TemporaryDirectory(prefix="deck-forge-pptx-") as td:
        slides = extract(html, a.browser_executable, a.image_only, Path(td))
        n = build_pptx(slides, Path(a.out), a.font)
    kinds = {}
    for s in slides:
        for it in s["items"]:
            kinds[it["kind"]] = kinds.get(it["kind"], 0) + 1
    print(f"  OK  {a.out}  ({n} slides; " + ", ".join(f"{v} {k}" for k, v in sorted(kinds.items())) + ")")
    if not a.image_only:
        print("  Next: render it with scripts/render_pptx.ps1 and run "
              "`export_pptx.py diff <ref_png_dir> <render_dir>`; see references/html-to-pptx.md.")


# ---------------------------------------------------------------------------
# diff
# ---------------------------------------------------------------------------
def cmd_diff(a):
    import numpy as np
    from PIL import Image, ImageFilter
    ref_dir, tgt_dir = Path(a.ref_dir), Path(a.target_dir)
    refs = sorted(ref_dir.glob("slide-*.png"))
    if not refs:
        raise SystemExit(f"no slide-NNN.png in {ref_dir}")
    report, worst = [], 0.0
    for r in refs:
        t = tgt_dir / r.name
        if not t.is_file():
            report.append({"page": r.name, "missing": True}); continue
        ra = Image.open(r).convert("RGB")
        ta = Image.open(t).convert("RGB")
        if ta.size != ra.size:
            ta = ta.resize(ra.size, Image.LANCZOS)
        # per-pixel L1 distance over RGB, thresholded
        d = (np.abs(np.asarray(ra, dtype=int) - np.asarray(ta, dtype=int)).sum(axis=2) > a.threshold)
        raw = float(d.mean() * 100)
        mask = Image.fromarray((d * 255).astype("uint8")).filter(ImageFilter.MinFilter(a.erode))
        solid = float((np.asarray(mask) > 0).mean() * 100)
        worst = max(worst, solid)
        report.append({"page": r.name, "raw_pct": round(raw, 2), "solid_pct": round(solid, 3)})
    if a.json:
        print(json.dumps({"pages": report, "max_solid_pct": round(worst, 3),
                          "max_solid_allowed": a.max_solid}, ensure_ascii=False, indent=1))
    else:
        # keep console output ASCII: Windows consoles decode child output as GBK
        print(f"{'page':<14}{'raw diff':>10}{'solid diff':>12}   (raw ~ rasterisation; solid = layout)")
        for row in report:
            if row.get("missing"):
                print(f"{row['page']:<14}{'MISSING':>10}"); continue
            flag = "  <-- inspect" if row["solid_pct"] > a.max_solid else ""
            print(f"{row['page']:<14}{row['raw_pct']:>9.2f}%{row['solid_pct']:>11.3f}%{flag}")
        print(f"max solid diff {worst:.3f}% (allowed {a.max_solid}%)")
    bad = any(r.get("missing") or r.get("solid_pct", 0) > a.max_solid for r in report)
    sys.exit(1 if bad else 0)


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    b = sub.add_parser("build", help="HTML deck -> editable .pptx")
    b.add_argument("html"); b.add_argument("out")
    b.add_argument("--font", default="Microsoft YaHei",
                   help="typeface written to every run (latin + east-asian); default is Windows-safe")
    b.add_argument("--image-only", action="store_true", help="one full-bleed screenshot per slide instead of shapes")
    b.add_argument("--browser-executable", default=None, help="reuse a local chrome.exe (see SKILL.md)")
    b.set_defaults(fn=cmd_build)
    d = sub.add_parser("diff", help="compare deck page renders with PPTX renders")
    d.add_argument("ref_dir"); d.add_argument("target_dir")
    d.add_argument("--erode", type=int, default=9, help="MinFilter size; strokes thinner than this are ignored")
    d.add_argument("--threshold", type=int, default=60, help="RGB L1 distance that counts as a changed pixel")
    d.add_argument("--max-solid", type=float, default=1.5, help="fail if any page's solid diff exceeds this %%")
    d.add_argument("--json", action="store_true")
    d.set_defaults(fn=cmd_diff)
    a = ap.parse_args(argv)
    a.fn(a)


if __name__ == "__main__":
    main()
