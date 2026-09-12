#!/usr/bin/env python3
"""Regression tests for scripts/export_pptx.py (synthetic deck, no renderer needed)."""

from __future__ import annotations

import subprocess
import sys
import tempfile
import unittest
from pathlib import Path


ROOT = Path(__file__).resolve().parent.parent
SCRIPT = ROOT / "scripts" / "export_pptx.py"
NS_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"

# One slide that exercises every emitter: a template edge bar (::after), a
# rounded box with an inline bold/red run, a CSS-border triangle under
# box-sizing:border-box, a table, an SVG mark, and a full-width centred line
# whose measurement slack would poke off-canvas without clamping.
DECK = """<!doctype html><html><head><meta charset="utf-8"><style>
*{box-sizing:border-box;margin:0}
html,body{width:100%;height:100%;overflow:hidden}
.deck-stage{position:absolute;left:0;top:0;width:1920px;height:1080px;overflow:hidden}
.slide{position:absolute;inset:0;width:1920px;height:1080px;overflow:hidden;visibility:hidden;background:#fff;font-family:sans-serif}
.slide.active,.slide.visible{visibility:visible}
.slide::after{content:'';position:absolute;top:0;right:0;width:38px;height:1080px;background:#0E106A}
.box{position:absolute;left:100px;top:200px;width:600px;height:200px;border:2.5px solid #000;border-radius:14px;background:#FFCCFF;padding:20px;font-size:29px;line-height:1.5}
.red{color:#f00;font-weight:700}
.tri{position:absolute;left:800px;top:280px;width:0;height:0;border-top:24px solid transparent;border-bottom:24px solid transparent;border-left:38px solid #B7B7C2}
table{position:absolute;left:100px;top:500px;width:600px;border-collapse:collapse}
td,th{border:2px solid #000;padding:10px;font-size:27px}
th{background:#F1F1F4}
.foot{position:absolute;left:0;top:1020px;width:1920px;text-align:center;font-size:21px}
.reveal{opacity:0;transform:translateY(22px)}
</style></head><body><main class="deck-stage">
<section class="slide active">
  <div class="box reveal">经济品并非<b>同一件少收钱</b>，<span class="red">低价写进产品本身</span>。</div>
  <div class="tri"></div>
  <table><tr><th>指标</th><th>2018</th></tr><tr><td>收入</td><td>6.25 亿元</td></tr></table>
  <svg style="position:absolute;left:60px;top:30px" width="132" height="132" viewBox="0 0 132 132"><ellipse cx="66" cy="60" rx="20" ry="40" fill="none" stroke="#000" stroke-width="3"/></svg>
  <div class="foot">-1-</div>
</section>
</main></body></html>"""


def run(*args: str) -> subprocess.CompletedProcess:
    return subprocess.run([sys.executable, str(SCRIPT), *args], capture_output=True, text=True, encoding="utf-8")


class ExportPptxTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        try:
            import playwright  # noqa: F401
            import pptx  # noqa: F401
        except ImportError as exc:  # pragma: no cover - environment dependent
            raise unittest.SkipTest(f"dependency missing: {exc}")

    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory(prefix="deckforge-export-pptx-")
        self.temp = Path(self.temp_dir.name)
        self.html = self.temp / "index.html"
        self.html.write_text(DECK, encoding="utf-8")

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def test_build_emits_native_shapes_inside_canvas(self) -> None:
        out = self.temp / "deck.pptx"
        res = run("build", str(self.html), str(out), "--font", "Test Sans")
        self.assertEqual(res.returncode, 0, res.stdout + res.stderr)
        self.assertIn("1 tri", res.stdout)  # the CSS triangle survived box-sizing:border-box

        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        prs = Presentation(str(out))
        self.assertEqual(len(prs.slides), 1)
        slide = prs.slides[0]
        kinds = [sh.shape_type for sh in slide.shapes]
        self.assertIn(MSO_SHAPE_TYPE.PICTURE, kinds, "svg should become a picture")
        self.assertIn(MSO_SHAPE_TYPE.TABLE, kinds)
        self.assertIn(MSO_SHAPE_TYPE.AUTO_SHAPE, kinds)

        W, H = prs.slide_width, prs.slide_height
        for sh in slide.shapes:
            self.assertGreaterEqual(sh.left, 0, f"{sh.name} pokes off the left edge")
            self.assertGreaterEqual(sh.top, 0)
            self.assertLessEqual(sh.left + sh.width, W, f"{sh.name} pokes off the right edge")
            self.assertLessEqual(sh.top + sh.height, H)

        # the ::after edge bar was emitted, flush right
        bars = [sh for sh in slide.shapes
                if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and sh.left + sh.width == W and sh.height == H]
        self.assertEqual(len(bars), 1, "template edge bar (::after) missing")

        texts = [sh for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        self.assertTrue(texts)
        runs = [r for sh in texts for p in sh.text_frame.paragraphs for r in p.runs]
        self.assertTrue(runs)
        for r in runs:
            rPr = r._r.find(NS_A + "rPr")
            self.assertEqual(rPr.get("lang"), "zh-CN", "runs must declare a CJK lang for 禁则")
            self.assertEqual(rPr.find(NS_A + "ea").get("typeface"), "Test Sans")
        self.assertTrue(any(r.font.bold and r.text == "同一件少收钱" for r in runs))
        red = [r for r in runs if r.text == "低价写进产品本身"]
        self.assertEqual(str(red[0].font.color.rgb), "FF0000")
        self.assertTrue(all(sh.text_frame._txBody.find(NS_A + "bodyPr").find(NS_A + "noAutofit") is not None
                            for sh in texts), "autofit must be disabled")

    def test_image_only_places_one_full_bleed_picture(self) -> None:
        out = self.temp / "img.pptx"
        res = run("build", str(self.html), str(out), "--image-only")
        self.assertEqual(res.returncode, 0, res.stdout + res.stderr)
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        prs = Presentation(str(out))
        shapes = list(prs.slides[0].shapes)
        self.assertEqual(len(shapes), 1)
        self.assertEqual(shapes[0].shape_type, MSO_SHAPE_TYPE.PICTURE)
        self.assertEqual((shapes[0].width, shapes[0].height), (prs.slide_width, prs.slide_height))

    def test_diff_separates_rasterisation_from_layout(self) -> None:
        from PIL import Image, ImageDraw

        ref, tgt = self.temp / "ref", self.temp / "tgt"
        ref.mkdir(); tgt.mkdir()
        base = Image.new("RGB", (400, 200), "white")
        ImageDraw.Draw(base).text((20, 20), "hello", fill="black")
        base.save(ref / "slide-001.png")
        # thin difference only: shift the text one pixel → rasterisation-class noise
        noisy = Image.new("RGB", (400, 200), "white")
        ImageDraw.Draw(noisy).text((21, 20), "hello", fill="black")
        noisy.save(tgt / "slide-001.png")
        res = run("diff", str(ref), str(tgt))
        self.assertEqual(res.returncode, 0, res.stdout + res.stderr)
        self.assertIn("0.000%", res.stdout)
        # a solid block that survives erosion → layout defect → exit 1
        block = noisy.copy()
        ImageDraw.Draw(block).rectangle((100, 60, 300, 160), fill="black")
        block.save(tgt / "slide-001.png")
        res = run("diff", str(ref), str(tgt))
        self.assertEqual(res.returncode, 1, res.stdout + res.stderr)
        self.assertIn("<-- inspect", res.stdout)


if __name__ == "__main__":
    unittest.main()
